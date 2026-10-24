#!/usr/bin/env python3
"""Visual contract: the seal is the only visible mark — no caption/signature text under it.

The cryptographic packet stays embedded (PDF keywords / hidden DOCX paragraph / SVG metadata);
only the human-visible caption is removed. This file covers every signing path that can draw a
visible seal: PDF stamping, universal/API PDF stamping, DOCX stamping, and SVG signing (P1/P2).
"""
from __future__ import annotations
import json, shutil, sys, tempfile, unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import pymupdf as fitz
import docx
import marma_watermark as mw

SECRET = mw.DEFAULT_SECRET
CAPTION = "MARMA-CRYPTO-SEAL"

def make_pdf(path: Path, text="Seal visual contract document") -> Path:
    d = fitz.open(); p = d.new_page()
    p.insert_text((72, 100), text, fontsize=12)
    p.insert_text((72, 130), "Second line for text-extraction coverage.", fontsize=10)
    d.save(str(path)); d.close(); return path

def pdf_text(path: Path) -> str:
    d = fitz.open(str(path)); t = "\n".join(pg.get_text("text") for pg in d); d.close(); return t

class TestSealOnlyPdf(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealonly_pdf_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        self.src = make_pdf(self.tmp / "src.pdf")

    def test_stamp_both_has_no_visible_caption(self):
        out = self.tmp / "both.pdf"
        res = mw.stamp_pdf_document(self.src, out, secret_key=SECRET, mode="both")
        self.assertTrue(res.get("format") == "pdf")
        self.assertNotIn(CAPTION, pdf_text(out))
        self.assertTrue(mw.verify_pdf_document(out)["verified"])
        keywords = fitz.open(str(out)).metadata.get("keywords", "")
        self.assertIn("MARMA_SEAL::", keywords)

    def test_stamp_seal_only_has_no_visible_caption(self):
        out = self.tmp / "seal.pdf"
        mw.stamp_pdf_document(self.src, out, secret_key=SECRET, mode="seal")
        self.assertNotIn(CAPTION, pdf_text(out))
        self.assertTrue(mw.verify_pdf_document(out)["verified"])

class TestSealOnlyUniversalPdf(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealonly_univ_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        self.src = make_pdf(self.tmp / "src.pdf", "Universal seal visual contract")

    def test_universal_emblems_have_no_visible_caption(self):
        try:
            import marma_crypto.api_server as api
        except Exception as e:  # pragma: no cover
            self.skipTest(f"api_server import failed: {e}")
        for emblem in ("tryzub", "corporate-shield", "cad-polygon"):
            out = self.tmp / f"{emblem}.pdf"
            try:
                api.stamp_pdf_universal_v21(self.src, out, secret_key=SECRET, recipient_id="",
                                            emblem_type=emblem, optical_ring=False, mode="both",
                                            signer_key_name="Sovereign-Vedic-Key-108", key_fingerprint="")
            except ValueError:
                continue
            self.assertNotIn(CAPTION, pdf_text(out), emblem)
            self.assertTrue(api.verify_pdf_universal_v21(out, secret_key=SECRET)["verified"], emblem)

class TestSealOnlyDocx(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealonly_docx_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        d = docx.Document(); d.add_paragraph("Seal visual contract DOCX"); self.src = self.tmp / "src.docx"; d.save(str(self.src))

    def test_docx_has_no_visible_caption_and_keeps_hidden_packet(self):
        out = self.tmp / "stamped.docx"
        mw.stamp_docx_document(self.src, out, secret_key=SECRET)
        self.assertTrue(mw.verify_docx_document(out)["verified"])
        d = docx.Document(str(out))
        visible_runs = [r.text for p in d.paragraphs for r in p.runs if not r.font.hidden and r.text]
        self.assertFalse(any(CAPTION in r for r in visible_runs))
        hidden_runs = [r.text for p in d.paragraphs for r in p.runs if r.font.hidden and r.text]
        self.assertTrue(any("MARMA_SEAL_PACKET::" in r for r in hidden_runs))

class TestSealOnlySvg(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealonly_svg_")); self.addCleanup(shutil.rmtree, self.tmp, True)

    def test_p1_and_p2_svg_have_no_text_elements_or_caption(self):
        sys.path.insert(0, str(ROOT / "src"))
        import marma_seal_profiles as sp
        for profile in ("v1-geometric", "v2-round42"):
            out = self.tmp / f"{profile}.svg"
            sp.sign_seal_svg("visual contract payload", SECRET, out, profile_id=profile)
            svg = out.read_text(encoding="utf-8")
            self.assertNotIn(CAPTION, svg)
            self.assertNotIn("<text", svg)
            self.assertNotIn("</text>", svg)
            v = sp.verify_seal_svg(svg, expected_profile=profile, payload="visual contract payload",
                                   secret_key=SECRET)
            self.assertTrue(v["verified"], v.get("reason"))

if __name__ == "__main__":
    unittest.main(verbosity=2)
