#!/usr/bin/env python3
"""Seal placement policy tests: last (default) | first | first-last | all.

Covers PDF, universal/API PDF, DOCX and SVG; verifies that the visible seal is present exactly
where the metadata says, and that a policy mismatch between metadata and visible layer fails.
"""
from __future__ import annotations
import json, os, shutil, subprocess, sys, tempfile, unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import pymupdf as fitz
import docx
import marma_watermark as mw

SECRET = mw.DEFAULT_SECRET
CLI = Path.home() / ".local" / "bin" / "marma-watermark"

def pdf3(path: Path) -> Path:
    d = fitz.open()
    for i in range(3):
        p = d.new_page(); p.insert_text((72, 100), f"Seal placement page {i+1}", fontsize=13)
    d.save(str(path)); d.close(); return path

def docx_doc(path: Path) -> Path:
    d = docx.Document()
    for i in range(5):
        d.add_paragraph(f"DOCX placement paragraph {i+1}")
    d.save(str(path)); return path

def packet_of(path: Path) -> dict:
    kw = fitz.open(str(path)).metadata.get("keywords", "")
    return json.loads(kw.split("MARMA_SEAL::", 1)[1])

def set_packet(path: Path, pkt: dict) -> None:
    d = fitz.open(str(path)); meta = d.metadata or {}
    meta["keywords"] = "MARMA_SEAL::" + json.dumps(pkt, separators=(",", ":"))
    d.set_metadata(meta); d.saveIncr(); d.close()

class TestResolve(unittest.TestCase):
    def test_policies(self):
        self.assertEqual(mw.resolve_seal_pages("last", 3), [2])
        self.assertEqual(mw.resolve_seal_pages("first", 3), [0])
        self.assertEqual(mw.resolve_seal_pages("first-last", 3), [0, 2])
        self.assertEqual(mw.resolve_seal_pages("all", 3), [0, 1, 2])
        self.assertEqual(mw.resolve_seal_pages("first-last", 1), [0])
        self.assertEqual(mw.resolve_seal_pages("all", 1), [0])

    def test_invalid_and_env(self):
        with self.assertRaises(ValueError):
            mw.resolve_seal_pages("middle", 3)
        old = os.environ.get("MARMA_SEAL_PAGES")
        try:
            os.environ["MARMA_SEAL_PAGES"] = "first"
            self.assertEqual(mw.resolve_seal_pages(None, 3), [0])
        finally:
            if old is None: os.environ.pop("MARMA_SEAL_PAGES", None)
            else: os.environ["MARMA_SEAL_PAGES"] = old

class TestPdfPlacement(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealpages_pdf_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        self.src = pdf3(self.tmp / "src.pdf")

    def test_all_policies(self):
        for spec in ("last", "first", "first-last", "all"):
            out = self.tmp / f"{spec}.pdf"
            mw.stamp_pdf_document(self.src, out, mode="both", seal_pages=spec)
            v = mw.verify_pdf_document(out)
            self.assertTrue(v["verified"], (spec, v.get("reason")))
            self.assertEqual(v["seal_pages"], spec)
            self.assertTrue(v["visible_seal_ok"])
            doc = fitz.open(str(out))
            expected = set(mw.resolve_seal_pages(spec, 3))
            actual = {i for i in range(3) if mw.page_has_visible_seal(doc[i])}
            doc.close()
            self.assertEqual(actual, expected, spec)

    def test_default_is_last(self):
        out = self.tmp / "default.pdf"; mw.stamp_pdf_document(self.src, out)
        self.assertEqual(packet_of(out)["seal_pages"], "last")
        v = mw.verify_pdf_document(out); self.assertTrue(v["verified"])

    def test_env_fallback(self):
        out = self.tmp / "env.pdf"; os.environ["MARMA_SEAL_PAGES"] = "first"
        try: mw.stamp_pdf_document(self.src, out)
        finally: os.environ.pop("MARMA_SEAL_PAGES", None)
        self.assertEqual(packet_of(out)["seal_pages"], "first")
        self.assertTrue(mw.verify_pdf_document(out)["verified"])

    def test_tampered_policy_fails_visible(self):
        out = self.tmp / "last.pdf"; mw.stamp_pdf_document(self.src, out, seal_pages="last")
        pkt = packet_of(out); pkt["seal_pages"] = "all"; set_packet(out, pkt)
        v = mw.verify_pdf_document(out)
        self.assertFalse(v["verified"])
        self.assertIn("Visible seal missing", v["reason"])
        self.assertEqual(sorted(v["visible_seal_missing_pages"]), [0, 1])

    def test_legacy_packet_without_policy_still_verifies(self):
        out = self.tmp / "legacy.pdf"; mw.stamp_pdf_document(self.src, out, seal_pages="first-last")
        pkt = packet_of(out); pkt.pop("seal_pages", None); set_packet(out, pkt)
        v = mw.verify_pdf_document(out)
        self.assertTrue(v["verified"], v.get("reason"))
        self.assertIsNone(v["visible_seal_ok"])

class TestDocxPlacement(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealpages_docx_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        self.src = docx_doc(self.tmp / "src.docx")

    def test_all_policies(self):
        for spec in ("last", "first", "first-last", "all"):
            out = self.tmp / f"{spec}.docx"
            mw.stamp_docx_document(self.src, out, seal_pages=spec)
            v = mw.verify_docx_document(out)
            self.assertTrue(v["verified"], (spec, v.get("reason")))
            self.assertEqual(v["seal_pages"], spec)
            self.assertTrue(v["visible_seal_ok"])

    def test_tampered_docx_policy_fails(self):
        out = self.tmp / "last.docx"; mw.stamp_docx_document(self.src, out, seal_pages="last")
        d = docx.Document(str(out))
        hidden = [p for p in d.paragraphs if "MARMA_SEAL_PACKET::" in p.text][0]
        pkt = json.loads(hidden.text.split("MARMA_SEAL_PACKET::", 1)[1]); pkt["seal_pages"] = "all"
        hidden.runs[0].text = "MARMA_SEAL_PACKET::" + json.dumps(pkt, separators=(",", ":"))
        hidden.runs[0].font.hidden = True
        d.save(str(out))
        v = mw.verify_docx_document(out)
        self.assertFalse(v["verified"])
        self.assertIn("Visible seal missing", v["reason"])

class TestUniversalAndSvgPlacement(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealpages_univ_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        self.src = pdf3(self.tmp / "src.pdf")

    def test_universal_pdf_policy(self):
        from marma_crypto.api_server import stamp_pdf_universal_v21, verify_pdf_universal_v21
        for spec in ("last", "all"):
            out = self.tmp / f"univ_{spec}.pdf"
            stamp_pdf_universal_v21(self.src, out, secret_key=SECRET, emblem_type="tryzub",
                                    optical_ring=False, seal_pages=spec)
            v = verify_pdf_universal_v21(out, secret_key=SECRET)
            self.assertTrue(v["verified"], (spec, v.get("reason")))
            self.assertEqual(v["seal_pages"], spec)
            self.assertTrue(v["visible_seal_ok"])

    def test_svg_seal_pages_metadata(self):
        out = self.tmp / "seal.svg"
        mw.sign_sri_yantra_svg("placement svg", SECRET, out)
        self.assertIn('data-seal-pages="last"', out.read_text(encoding="utf-8"))
        v = mw.verify_sri_yantra_svg(out.read_text(encoding="utf-8"), "placement svg", SECRET)
        self.assertTrue(v["verified"])
        self.assertEqual(v["seal_pages"], "last")

class TestCliSealPages(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sealpages_cli_")); self.addCleanup(shutil.rmtree, self.tmp, True)
        self.src = pdf3(self.tmp / "src.pdf")

    def _cli(self, *args):
        return subprocess.run([str(CLI), *args], cwd=str(self.tmp), capture_output=True, text=True, timeout=120)

    def _verify(self, path):
        r = self._cli("verify", str(path)); t = r.stdout
        return json.loads(t[t.find("{"):]) if "{" in t else {}

    def test_help_lists_choices(self):
        r = self._cli("--help")
        self.assertEqual(r.returncode, 0); self.assertIn("--seal-pages", r.stdout)
        for c in ("last", "first", "first-last", "all"): self.assertIn(c, r.stdout)

    def test_cli_policies_and_default(self):
        for spec in ("last", "first-last", "all"):
            out = self.tmp / f"cli_{spec}.pdf"
            r = self._cli("sign", "--mode", "seal", "--seal-pages", spec, "-o", str(out), str(self.src))
            self.assertEqual(r.returncode, 0, r.stderr[-400:])
            v = self._verify(out)
            self.assertTrue(v.get("verified"), (spec, v.get("reason")))
            self.assertEqual(v.get("seal_pages"), spec)
        out = self.tmp / "cli_default.pdf"
        self._cli("sign", "--mode", "seal", "-o", str(out), str(self.src))
        self.assertEqual(self._verify(out).get("seal_pages"), "last")

    def test_invalid_policy_rejected(self):
        r = self._cli("sign", "--seal-pages", "middle", str(self.src))
        self.assertNotEqual(r.returncode, 0)

if __name__ == "__main__":
    unittest.main(verbosity=2)
